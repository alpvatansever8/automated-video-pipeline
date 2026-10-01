import pytest
from docx import Document

from video_pipeline.script_parser import parse_script

TXT = """# First

Paragraph one.

Paragraph two is a bit longer than the first one.

# Second

Only paragraph.
"""


def test_txt_sections_and_paragraphs(tmp_path):
    script = tmp_path / "s.txt"
    script.write_text(TXT, encoding="utf-8")
    sections = parse_script(script, total_minutes=10, save_plan=False)
    assert [s["title"] for s in sections] == ["First", "Second"]
    assert len(sections[0]["paragraphs"]) == 2
    assert len(sections[1]["paragraphs"]) == 1


def test_durations_add_up_to_total(tmp_path):
    script = tmp_path / "s.txt"
    script.write_text(TXT, encoding="utf-8")
    sections = parse_script(script, total_minutes=10, save_plan=False)
    assert sum(s["duration_minutes"] for s in sections) == pytest.approx(10, abs=0.05)
    assert sections[0]["time_start_minutes"] == 0
    assert sections[1]["time_start_minutes"] == pytest.approx(sections[0]["time_end_minutes"])


def test_docx_bold_paragraph_is_a_title(tmp_path):
    doc = Document()
    doc.add_paragraph().add_run("Bold Title").bold = True
    doc.add_paragraph("Body text of the section.")
    path = tmp_path / "s.docx"
    doc.save(path)
    sections = parse_script(path, total_minutes=5, save_plan=False)
    assert sections[0]["title"] == "Bold Title"
    assert sections[0]["paragraphs"] == ["Body text of the section."]


def test_unsupported_format_is_rejected(tmp_path):
    bad = tmp_path / "s.pdf"
    bad.write_text("x")
    with pytest.raises(ValueError):
        parse_script(bad, total_minutes=5, save_plan=False)


def test_plan_file_is_written_to_output(tmp_path):
    script = tmp_path / "s.txt"
    script.write_text(TXT, encoding="utf-8")
    parse_script(script, total_minutes=5)
    assert (tmp_path / "output" / "s_plan.txt").exists()
